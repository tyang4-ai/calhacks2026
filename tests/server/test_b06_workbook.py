"""B06 · the month-close workbook. Contract tests for POST /api/b06/workbook and GET /api/b06/files.

The fixture b06_month.json is September's records, the scripted workbook.ready and the demo
figures, dumped from src/mock. openpyxl stores formulas without values, so `evaluate()` below
computes the few formula shapes the workbook uses (SUM, SUMIFS, COUNTA, cell arithmetic) to check
the totals against the fixture. Nothing here calls the network except the one test marked live,
which posts to the running API on localhost:8000 and skips when it doesn't list b06.
"""

from __future__ import annotations

import copy
import io
import json
import re
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook
from openpyxl.chart import BarChart
from openpyxl.utils import range_boundaries

from server.routers.b06_workbook import FILES_DIR, XLSX_TYPE

FIXTURE = Path(__file__).with_name("b06_month.json")
ENVELOPE = {"id", "type", "at", "source", "key", "payload"}
SHEETS = ["Ledger", "Reconciliation", "Profit and loss", "Makers"]
FILE_NAME = "cindia-close-2026-09.xlsx"


def load_month() -> dict:
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


def body_for(month: dict, trigger="fixture") -> dict:
    return {"trigger": month["trigger"] if trigger == "fixture" else trigger, "monthId": month["monthId"], "records": month["records"]}


# ---------------------------------------------------------------------------------------------
# A tiny evaluator for the workbook's formulas.


class Evaluator:
    def __init__(self, wb):
        self.wb = wb

    def cell(self, sheet: str, addr: str):
        v = self.wb[sheet][addr.replace("$", "")].value
        if isinstance(v, str) and v.startswith("="):
            return self.expr(sheet, v[1:])
        return v

    def values(self, sheet: str, ref: str) -> list:
        if "!" in ref:
            sheet, ref = ref.split("!")
        min_col, min_row, max_col, max_row = range_boundaries(ref.replace("$", ""))
        out = []
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                addr = f"{self.wb[sheet].cell(row=r, column=c).coordinate}"
                out.append(self.cell(sheet, addr))
        return out

    def expr(self, sheet: str, text: str) -> float:
        def sumifs(m):
            total = self.values(sheet, m.group(1))
            crit = self.values(sheet, m.group(2))
            return repr(sum(v or 0 for v, c in zip(total, crit) if c == m.group(3)))

        text = re.sub(r'SUMIFS\(([^,]+),([^,]+),"([^"]+)"\)', sumifs, text)
        text = re.sub(r"SUM\(([^)]+)\)", lambda m: repr(sum(v or 0 for v in self.values(sheet, m.group(1)))), text)
        text = re.sub(r"COUNTA\(([^)]+)\)", lambda m: repr(sum(1 for v in self.values(sheet, m.group(1)) if v not in (None, ""))), text)
        text = re.sub(r"\$?([A-Z]{1,2})\$?(\d+)", lambda m: repr(self.cell(sheet, m.group(1) + m.group(2)) or 0), text)
        assert re.fullmatch(r"[0-9.eE+\-*/() ]+", text), f"unexpected formula left: {text}"
        return eval(text)  # noqa: S307 - only digits and arithmetic remain


def cents(x: float) -> int:
    return round(x * 100)


def find_row(ws, col: str, label: str) -> int:
    for r in range(1, ws.max_row + 1):
        if ws[f"{col}{r}"].value == label:
            return r
    raise AssertionError(f"{label} not found in column {col} of {ws.title}")


@pytest.fixture
def built(client):
    month = load_month()
    res = client.post("/api/b06/workbook", json=body_for(month))
    assert res.status_code == 200, res.text
    got = client.get(f"/api/b06/files/{FILE_NAME}")
    assert got.status_code == 200
    wb = load_workbook(io.BytesIO(got.content))
    return month, res.json(), got, wb


# ---------------------------------------------------------------------------------------------


def test_response_shape(built):
    month, body, _got, _wb = built
    assert list(body) == ["events"]
    assert len(body["events"]) == 1
    ev = body["events"][0]
    assert set(ev) == ENVELOPE
    assert ev["type"] == "workbook.ready"
    assert ev["source"] == "live"
    assert ev["key"] == "b06-workbook"
    assert re.fullmatch(r"evt_live_\d+", ev["id"])
    assert ev["at"] == month["trigger"]["at"]
    assert ev["payload"] == {"monthId": "2026-09", "fileName": FILE_NAME, "sheets": SHEETS}
    assert any(FILES_DIR.glob("cindia-close-2026-09*.xlsx"))


def test_ids_are_unique_and_at_is_zero_without_trigger(client):
    month = load_month()
    a = client.post("/api/b06/workbook", json=body_for(month, trigger=None)).json()["events"][0]
    b = client.post("/api/b06/workbook", json=body_for(month, trigger=None)).json()["events"][0]
    assert a["at"] == 0
    assert a["id"] != b["id"]


def test_served_as_xlsx(built):
    _month, _body, got, _wb = built
    assert got.headers["content-type"] == XLSX_TYPE
    assert got.content[:2] == b"PK"


def test_sheet_names(built):
    *_, wb = built
    assert wb.sheetnames == SHEETS


def test_every_totals_row_is_a_formula(built):
    *_, wb = built
    ledger = wb["Ledger"]
    t = find_row(ledger, "A", "Total")
    assert ledger[f"D{t}"].value.startswith("=SUM(") and ledger[f"E{t}"].value.startswith("=SUM(")
    rec = wb["Reconciliation"]
    t = find_row(rec, "A", "Total")
    for col in "DEFGHI":
        assert rec[f"{col}{t}"].value.startswith("=SUM(")
    pnl = wb["Profit and loss"]
    for label in ("Net sales", "Total costs", "Net income"):
        assert str(pnl[f"B{find_row(pnl, 'A', label)}"].value).startswith("=")
    for r in range(2, find_row(pnl, "A", "Net income")):
        assert "Ledger!" in pnl[f"B{r}"].value or pnl[f"B{r}"].value.startswith("=SUM(")
    makers = wb["Makers"]
    t = find_row(makers, "A", "Total")
    for col in "DEF":
        assert makers[f"{col}{t}"].value.startswith("=SUM(")


def test_one_native_bar_chart_on_makers(built):
    *_, wb = built
    charts = wb["Makers"]._charts
    assert len(charts) == 1
    assert isinstance(charts[0], BarChart)
    for other in ("Ledger", "Reconciliation", "Profit and loss"):
        assert wb[other]._charts == []


def test_totals_match_the_dataset(built):
    month, _body, _got, wb = built
    exp = month["expected"]
    ev = Evaluator(wb)
    ledger_ws = wb["Ledger"]
    t = find_row(ledger_ws, "A", "Total")
    assert ledger_ws.max_row == len(month["records"]["ledger"]) + 2
    assert cents(ev.cell("Ledger", f"D{t}")) == exp["ledgerDebitCents"]
    assert cents(ev.cell("Ledger", f"E{t}")) == exp["ledgerCreditCents"]

    makers = wb["Makers"]
    for r in range(2, 2 + len(exp["makerTotals"])):
        mk = makers[f"A{r}"].value
        want = exp["makerTotals"][mk]
        assert cents(ev.cell("Makers", f"D{r}")) == want["totalCents"]
        assert cents(ev.cell("Makers", f"E{r}")) == want["heldCents"]
        assert cents(ev.cell("Makers", f"F{r}")) == want["totalCents"] - want["heldCents"]
    t = find_row(makers, "A", "Total")
    assert cents(ev.cell("Makers", f"F{t}")) == exp["makerPayCents"]
    assert cents(ev.cell("Makers", f"E{t}")) == exp["heldCents"]

    rec = wb["Reconciliation"]
    r = find_row(rec, "A", "Receivable from the gallery") + 2
    assert cents(ev.cell("Reconciliation", f"F{r}")) == exp["galleryDeltaCents"] == 34000
    statuses = {rec[f"A{i}"].value: rec[f"J{i}"].value for i in range(2, find_row(rec, "A", "Total"))}
    gallery = [p["id"] for p in month["records"]["payouts"] if p["channelId"] == "ch_gallery"]
    assert all(statuses[g] == "short_payment" for g in gallery)
    assert "carried_over" in statuses.values() and "bundled" in statuses.values()

    pnl = wb["Profit and loss"]
    ledger = month["records"]["ledger"]

    def net(account: str, side: str) -> int:
        dr = sum(e["debitCents"] for e in ledger if e["account"] == account)
        cr = sum(e["creditCents"] for e in ledger if e["account"] == account)
        return cr - dr if side == "credit" else dr - cr

    assert cents(ev.cell("Profit and loss", f"B{find_row(pnl, 'A', 'Net sales')}")) == exp["salesCents"]
    assert cents(ev.cell("Profit and loss", f"B{find_row(pnl, 'A', 'Maker pay')}")) == net("maker_pay", "debit")
    costs = sum(net(a, "debit") for a in ("card_fees", "platform_fees", "gallery_commission", "materials", "maker_pay"))
    assert cents(ev.cell("Profit and loss", f"B{find_row(pnl, 'A', 'Total costs')}")) == costs
    assert cents(ev.cell("Profit and loss", f"B{find_row(pnl, 'A', 'Net income')}")) == exp["salesCents"] - costs


@pytest.mark.parametrize(
    "name",
    ["..%2F..%2Fserver%2Fconfig.py", "..%5C..%5Cserver%5Cconfig.py", "..", "..%2Fcindia-close-2026-09.xlsx",
     "nothing-here.xlsx", "secret.txt"],
)
def test_path_tricks_give_not_found(client, name):
    res = client.get(f"/api/b06/files/{name}")
    assert res.status_code == 404
    assert res.json()["error"]["code"] == "not_found"
    assert set(res.json()["error"]) == {"code", "message"}


def test_bad_body_gives_bad_request(client):
    month = load_month()
    for broken in ({"monthId": "2026-09"}, {**body_for(month), "monthId": "September"}):
        res = client.post("/api/b06/workbook", json=broken)
        assert res.status_code == 400
        assert res.json()["error"]["code"] == "bad_request"
    body = copy.deepcopy(body_for(month))
    body["records"]["ledger"][0]["debitCents"] = "lots"
    res = client.post("/api/b06/workbook", json=body)
    assert res.status_code == 400
    assert res.json()["error"]["code"] == "bad_request"


def test_no_key_needed(client, monkeypatch):
    # B06 needs no key, so it never answers key_missing.
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("NESSIE_KEY", raising=False)
    res = client.post("/api/b06/workbook", json=body_for(load_month()))
    assert res.status_code == 200


def test_build_succeeds_while_the_served_file_is_open(client):
    # On Windows a rename onto a file another handle has open fails with PermissionError. The API
    # on :8000 may be serving the workbook while this process builds it, so a build must not care.
    month = load_month()
    first = client.post("/api/b06/workbook", json=body_for(month))
    assert first.status_code == 200, first.text
    with client.stream("GET", f"/api/b06/files/{FILE_NAME}") as held:
        assert held.status_code == 200
        open_files = [open(p, "rb") for p in FILES_DIR.glob("cindia-close-2026-09*.xlsx")]
        try:
            res = client.post("/api/b06/workbook", json=body_for(month))
            assert res.status_code == 200, res.text
            assert res.json()["events"][0]["payload"]["fileName"] == FILE_NAME
            got = client.get(f"/api/b06/files/{FILE_NAME}")
            assert got.status_code == 200
            assert got.content[:2] == b"PK"
            assert load_workbook(io.BytesIO(got.content)).sheetnames == SHEETS
        finally:
            for f in open_files:
                f.close()


def test_download_keeps_the_contract_name(built):
    _month, _body, got, _wb = built
    assert FILE_NAME in got.headers["content-disposition"]


def test_listed_in_health(client):
    assert "b06" in client.get("/api/health").json()["routers"]


@pytest.mark.live
def test_live_running_api():
    """One real call to the API the orchestrator runs on :8000. Skips when it isn't up or lacks b06."""
    try:
        health = httpx.get("http://127.0.0.1:8000/api/health", timeout=2).json()
    except Exception:
        pytest.skip("the API on :8000 isn't running")
    if "b06" not in health.get("routers", []):
        pytest.skip("the running API hasn't loaded b06 yet (needs a restart)")
    res = httpx.post("http://127.0.0.1:8000/api/b06/workbook", json=body_for(load_month()), timeout=7)
    assert res.status_code == 200
    assert res.json()["events"][0]["payload"]["fileName"] == FILE_NAME
