"""B06 · the month-close workbook (`b06-workbook`).

POST /api/b06/workbook builds `cindia-close-<monthId>.xlsx` with openpyxl from the month's records
and returns one `workbook.ready`. GET /api/b06/files/{fileName} serves the newest build of that
workbook from artifacts/files/, downloaded under the contract name. Each build is saved under its
own name on disk (see save_workbook), so concurrent builds and downloads never collide on Windows.
No key, no network.

The sheets follow the browser's SheetJS workbook (src/features/close/buildWorkbook.ts), plus a
native bar chart on Makers, which SheetJS can't write:
- Ledger: one row per ledger entry, debit and credit totals as =SUM(...).
- Reconciliation: every channel payout with its orders, the bank deposit that carried it, a delta
  formula and its status, then the gallery's receivable: expected net at the contract's commission
  (the orders' commissionCents) against the bank deposit, the difference as a formula.
- Profit and loss: every line is a SUMIFS formula over the Ledger sheet, by account, with SUM totals.
- Makers: one summary row per maker (pieces, share, held, payout) with formulas over the maker's
  lines below it, and a clustered bar chart of share and payout per maker.

Amounts go in cells as dollars (cents / 100) with a money number format. openpyxl writes formulas
without cached values, so the workbook asks Excel to recalculate fully when it opens.
"""

from __future__ import annotations

import itertools
import os
import re
import secrets
import tempfile
import threading
import time
from pathlib import Path
from typing import Any, Optional

from fastapi import APIRouter
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.workbook.properties import CalcProperties
from pydantic import BaseModel, ConfigDict, Field

from server import ApiError

router = APIRouter(prefix="/api/b06")

KEY = "b06-workbook"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEETS = ["Ledger", "Reconciliation", "Profit and loss", "Makers"]
FILES_DIR = Path(__file__).resolve().parents[2] / "artifacts" / "files"
SAFE_NAME = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]*\.xlsx$")
KEEP = 5  # builds of one workbook kept on disk; older ones are removed when not open

MONEY = "#,##0.00"
RATE = "0%"
BOLD = Font(bold=True)

# evt_live_ plus a number, unique within the server process. B06's numbers start at 600001 so
# they never meet another router's counter.
_ids = itertools.count(600001)
_ids_lock = threading.Lock()
_write_lock = threading.Lock()


def _next_id() -> str:
    with _ids_lock:
        return f"evt_live_{next(_ids)}"


class Records(BaseModel):
    model_config = ConfigDict(extra="allow")

    orders: list[dict[str, Any]]
    payouts: list[dict[str, Any]]
    statements: list[dict[str, Any]]
    bankTxns: list[dict[str, Any]]
    ledger: list[dict[str, Any]]
    makerStatements: list[dict[str, Any]]


class WorkbookRequest(BaseModel):
    trigger: Optional[dict[str, Any]]
    monthId: str = Field(pattern=r"^\d{4}-\d{2}$")
    records: Records


def file_name_for(month_id: str) -> str:
    return f"cindia-close-{month_id}.xlsx"


# ---------------------------------------------------------------------------------------------
# Reading records. Every amount is integer cents; a bad one is the caller's mistake (bad_request).


def _cents(rec: dict[str, Any], field: str) -> int:
    v = rec.get(field, 0)
    if isinstance(v, bool) or not isinstance(v, int):
        raise ApiError("bad_request", f"The field {field} must be an integer number of cents.")
    return v


def _str(rec: dict[str, Any], field: str) -> str:
    v = rec.get(field, "")
    return v if isinstance(v, str) else str(v)


def _dollars(cents: int) -> float:
    return cents / 100


def _header(ws, row: int, labels: list[str]) -> None:
    for c, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font = BOLD


def _money(ws, addr: str, value: Any) -> None:
    ws[addr] = value
    ws[addr].number_format = MONEY


# ---------------------------------------------------------------------------------------------
# Ledger


def _ledger_sheet(ws, ledger: list[dict[str, Any]]) -> int:
    """Writes the Ledger sheet and returns the last entry row (header is row 1)."""
    _header(ws, 1, ["Entry", "Date", "Account", "Debit", "Credit", "Record", "Memo"])
    for i, e in enumerate(ledger):
        row = i + 2
        ws[f"A{row}"] = _str(e, "id")
        ws[f"B{row}"] = _str(e, "date")
        ws[f"C{row}"] = _str(e, "account")
        _money(ws, f"D{row}", _dollars(_cents(e, "debitCents")))
        _money(ws, f"E{row}", _dollars(_cents(e, "creditCents")))
        ws[f"F{row}"] = _str(e, "refId")
        ws[f"G{row}"] = _str(e, "memo")
    last = len(ledger) + 1
    total = last + 1
    ws[f"A{total}"] = "Total"
    ws[f"A{total}"].font = BOLD
    for col in ("D", "E"):
        _money(ws, f"{col}{total}", f"=SUM({col}2:{col}{max(last, 2)})")
        ws[f"{col}{total}"].font = BOLD
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["G"].width = 48
    return last


# ---------------------------------------------------------------------------------------------
# Reconciliation


def _status(p: dict[str, Any], month_id: str, deposit: int, expected_net: int) -> str:
    if _str(p, "monthId") and _str(p, "monthId") != month_id:
        return "carried_over"
    if deposit < expected_net and deposit > 0:
        return "short_payment"
    if _str(p, "channelId") == "ch_gallery":
        return "one_to_one" if deposit >= expected_net else "awaiting_deposit"
    if p.get("refundIds"):
        return "refund_netted"
    if len(p.get("orderIds") or []) > 1:
        return "bundled"
    if deposit == 0:
        return "awaiting_deposit"
    return "one_to_one"


def _reconciliation_sheet(ws, records: Records, month_id: str) -> None:
    deposits: dict[str, int] = {}
    for t in records.bankTxns:
        if _str(t, "direction") == "in":
            ref = _str(t, "refId")
            deposits[ref] = deposits.get(ref, 0) + _cents(t, "amountCents")

    orders = {_str(o, "id"): o for o in records.orders}
    gallery_orders = [o for o in records.orders if _str(o, "channelId") == "ch_gallery"]
    gallery_payouts = [p for p in records.payouts if _str(p, "channelId") == "ch_gallery"]

    _header(ws, 1, ["Payout", "Channel", "Date", "Orders gross", "Fees", "Refund", "Expected net",
                    "Bank deposit", "Delta", "Status", "Orders", "Month"])
    payouts = sorted(records.payouts, key=lambda p: (_str(p, "date"), _str(p, "id")))
    for i, p in enumerate(payouts):
        row = i + 2
        pid = _str(p, "id")
        deposit = deposits.get(pid, 0)
        gross = _cents(p, "grossCents")
        fees = _cents(p, "feesCents")
        if _str(p, "channelId") == "ch_gallery":
            # The gallery's expected net is the contract's: gross less the orders' booked commission.
            fees = sum(_cents(orders[o], "commissionCents") for o in (p.get("orderIds") or []) if o in orders)
        ws[f"A{row}"] = pid
        ws[f"B{row}"] = _str(p, "channelId")
        ws[f"C{row}"] = _str(p, "date")
        _money(ws, f"D{row}", _dollars(gross))
        _money(ws, f"E{row}", _dollars(fees))
        _money(ws, f"F{row}", _dollars(_cents(p, "refundCents")))
        _money(ws, f"G{row}", f"=D{row}-E{row}-F{row}")
        _money(ws, f"H{row}", _dollars(deposit))
        _money(ws, f"I{row}", f"=G{row}-H{row}")
        ws[f"J{row}"] = _status(p, month_id, deposit, gross - fees - _cents(p, "refundCents"))
        ws[f"K{row}"] = ", ".join(str(o) for o in (p.get("orderIds") or []))
        ws[f"L{row}"] = _str(p, "monthId")
    first, last = 2, max(len(payouts) + 1, 2)
    total = last + 1
    ws[f"A{total}"] = "Total"
    ws[f"A{total}"].font = BOLD
    for col in "DEFGHI":
        _money(ws, f"{col}{total}", f"=SUM({col}{first}:{col}{last})")
        ws[f"{col}{total}"].font = BOLD

    # The gallery receivable: what the contract says it owes, against what reached the bank.
    top = total + 2
    ws[f"A{top}"] = "Receivable from the gallery"
    ws[f"A{top}"].font = BOLD
    _header(ws, top + 1, ["Statement", "Gross", "Contract rate", "Expected net", "Actual (bank)", "Receivable", "Records"])
    gross = sum(_cents(o, "grossCents") for o in gallery_orders)
    commission = sum(_cents(o, "commissionCents") for o in gallery_orders)
    actual = sum(deposits.get(_str(p, "id"), 0) for p in gallery_payouts)
    stm = next((s for s in records.statements if _str(s, "channelId") == "ch_gallery"), {})
    row = top + 2
    ws[f"A{row}"] = f"{_str(stm, 'id')} ({month_id})" if stm else month_id
    _money(ws, f"B{row}", _dollars(gross))
    ws[f"C{row}"] = (commission / gross) if gross else 0
    ws[f"C{row}"].number_format = RATE
    _money(ws, f"D{row}", f"=B{row}*(1-C{row})")
    _money(ws, f"E{row}", _dollars(actual))
    _money(ws, f"F{row}", f"=D{row}-E{row}")
    ws[f"F{row}"].font = BOLD
    ws[f"G{row}"] = ", ".join([_str(stm, "id")] + [_str(p, "id") for p in gallery_payouts]) if stm else ""
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["K"].width = 40


# ---------------------------------------------------------------------------------------------
# Profit and loss: SUMIFS over the Ledger sheet, by account.

SALES = [("Sales, Cindia", "sales_cindia", "credit"), ("Sales, co-op site", "sales_site", "credit"),
         ("Sales, gallery", "sales_gallery", "credit"), ("Refunds", "refunds", "refund")]
COSTS = [("Card fees", "card_fees", "debit"), ("Cindia fees", "platform_fees", "debit"),
         ("Gallery commission", "gallery_commission", "debit"), ("Materials", "materials", "debit"),
         ("Maker pay", "maker_pay", "debit")]


def _sumifs(col: str, account: str, last: int) -> str:
    return f'SUMIFS(Ledger!${col}$2:${col}${last},Ledger!$C$2:$C${last},"{account}")'


def _account_formula(account: str, side: str, last: int) -> str:
    dr, cr = _sumifs("D", account, last), _sumifs("E", account, last)
    if side == "credit":
        return f"={cr}-{dr}"
    if side == "refund":
        return f"=-({dr}-{cr})"
    return f"={dr}-{cr}"


def _pnl_sheet(ws, month_id: str, ledger_last: int) -> None:
    last = max(ledger_last, 2)
    ws["A1"] = f"Profit and loss, {month_id}"
    ws["B1"] = "Amount"
    ws["A1"].font = BOLD
    ws["B1"].font = BOLD
    row = 2
    for label, account, side in SALES:
        ws[f"A{row}"] = label
        _money(ws, f"B{row}", _account_formula(account, side, last))
        row += 1
    net_sales = row
    ws[f"A{row}"] = "Net sales"
    _money(ws, f"B{row}", f"=SUM(B2:B{row - 1})")
    row += 1
    first_cost = row
    for label, account, side in COSTS:
        ws[f"A{row}"] = label
        _money(ws, f"B{row}", _account_formula(account, side, last))
        row += 1
    total_costs = row
    ws[f"A{row}"] = "Total costs"
    _money(ws, f"B{row}", f"=SUM(B{first_cost}:B{row - 1})")
    row += 1
    ws[f"A{row}"] = "Net income"
    _money(ws, f"B{row}", f"=B{net_sales}-B{total_costs}")
    for r in (net_sales, total_costs, row):
        ws[f"A{r}"].font = BOLD
        ws[f"B{r}"].font = BOLD
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 14


# ---------------------------------------------------------------------------------------------
# Makers: a summary table with a bar chart, then each maker's lines.


def _makers_sheet(ws, statements: list[dict[str, Any]]) -> None:
    n = len(statements)
    _header(ws, 1, ["Maker", "Statement", "Pieces", "Maker share", "Held", "Payout", "Payout date"])
    summary_total = n + 2
    detail = summary_total + 3  # header row of the lines table
    _header(ws, detail, ["Maker", "Order", "Piece", "Maker share"])
    row = detail + 1
    spans: list[tuple[int, int]] = []
    for st in statements:
        lines = st.get("lines") or []
        first = row
        for line in lines:
            if not isinstance(line, dict):
                raise ApiError("bad_request", "Each maker statement line must be an object.")
            ws[f"A{row}"] = _str(st, "makerId")
            ws[f"B{row}"] = _str(line, "orderId")
            ws[f"C{row}"] = _str(line, "pieceId")
            _money(ws, f"D{row}", _dollars(_cents(line, "makerShareCents")))
            row += 1
        spans.append((first, row - 1))
    for i, (st, (first, last)) in enumerate(zip(statements, spans)):
        r = i + 2
        ws[f"A{r}"] = _str(st, "makerId")
        ws[f"B{r}"] = _str(st, "id")
        if last >= first:
            ws[f"C{r}"] = f"=COUNTA(C{first}:C{last})"
            _money(ws, f"D{r}", f"=SUM(D{first}:D{last})")
        else:
            ws[f"C{r}"] = 0
            _money(ws, f"D{r}", 0)
        _money(ws, f"E{r}", _dollars(_cents(st, "heldCents")))
        _money(ws, f"F{r}", f"=D{r}-E{r}")
        ws[f"G{r}"] = _str(st, "payoutDate")
    ws[f"A{summary_total}"] = "Total"
    for col in "CDEF":
        cell = f"{col}{summary_total}"
        if col == "C":
            ws[cell] = f"=SUM(C2:C{max(n + 1, 2)})"
        else:
            _money(ws, cell, f"=SUM({col}2:{col}{max(n + 1, 2)})")
        ws[cell].font = BOLD
    ws[f"A{summary_total}"].font = BOLD

    if n:
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
        chart.title = "Maker earnings"
        chart.y_axis.title = "US dollars"
        chart.x_axis.title = "Maker"
        data = Reference(ws, min_col=4, max_col=4, min_row=1, max_row=n + 1)
        payout = Reference(ws, min_col=6, max_col=6, min_row=1, max_row=n + 1)
        chart.add_data(data, titles_from_data=True)
        chart.add_data(payout, titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
        chart.width = 16
        chart.height = 8
        ws.add_chart(chart, "I2")
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["D"].width = 14


# ---------------------------------------------------------------------------------------------


def build_workbook(records: Records, month_id: str) -> Workbook:
    """Builds the four-sheet workbook in memory. Pure: no file, no clock."""
    wb = Workbook()
    ledger_ws = wb.active
    ledger_ws.title = "Ledger"
    ledger_last = _ledger_sheet(ledger_ws, records.ledger)
    _reconciliation_sheet(wb.create_sheet("Reconciliation"), records, month_id)
    _pnl_sheet(wb.create_sheet("Profit and loss"), month_id, ledger_last)
    _makers_sheet(wb.create_sheet("Makers"), records.makerStatements)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    return wb


def _version_pattern(file_name: str) -> re.Pattern[str]:
    stem = file_name[: -len(".xlsx")]
    return re.compile(re.escape(stem) + r"--(\d{20})-[0-9a-f]{8}\.xlsx")


def _versions(file_name: str, folder: Path) -> list[Path]:
    """Every build of `file_name` in `folder`, oldest first."""
    pattern = _version_pattern(file_name)
    try:
        names = [p.name for p in folder.iterdir()]
    except OSError:
        return []
    return [folder / n for n in sorted(n for n in names if pattern.fullmatch(n))]


def save_workbook(wb: Workbook, file_name: str, folder: Path = FILES_DIR) -> Path:
    """Saves one build under its own name, `<stem>--<time_ns>-<random>.xlsx`, and returns its path.

    The name the contract gives (`cindia-close-2026-09.xlsx`) is resolved to the newest build by
    `resolve_file`. A fixed target name failed on Windows: renaming onto a file another process has
    open (the API on :8000 serving it, or a test reading it) raises PermissionError. A fresh name is
    never open, so a build can't collide with a download. The temp file plus rename means a download
    never sees a half-written workbook. Older builds beyond the newest KEEP are removed when nobody
    holds them open; one that's open is left for the next build to remove.
    """
    folder.mkdir(parents=True, exist_ok=True)
    stem = file_name[: -len(".xlsx")]
    target = folder / f"{stem}--{time.time_ns():020d}-{secrets.token_hex(4)}.xlsx"
    with _write_lock:
        fd, tmp = tempfile.mkstemp(suffix=".xlsx.tmp", dir=folder)
        os.close(fd)
        try:
            wb.save(tmp)
            os.replace(tmp, target)
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)
        for old in [p for p in _versions(file_name, folder) if p != target][: -(KEEP - 1) or None]:
            try:
                old.unlink()
            except OSError:
                pass
    return target


def resolve_file(file_name: str, folder: Path = FILES_DIR) -> Optional[Path]:
    """The newest build of `file_name`, else a file saved under exactly that name, else None."""
    builds = [p for p in _versions(file_name, folder) if p.is_file()]
    if builds:
        return builds[-1]
    exact = folder / file_name
    return exact if exact.is_file() else None


@router.post("/workbook")
def make_workbook(body: WorkbookRequest) -> dict:
    file_name = file_name_for(body.monthId)
    wb = build_workbook(body.records, body.monthId)
    try:
        save_workbook(wb, file_name)
    except OSError:
        raise ApiError("internal", "The server couldn't save the workbook.")
    at = body.trigger.get("at") if isinstance(body.trigger, dict) else None
    event = {
        "id": _next_id(),
        "type": "workbook.ready",
        "at": at if isinstance(at, (int, float)) and not isinstance(at, bool) else 0,
        "source": "live",
        "key": KEY,
        "payload": {"monthId": body.monthId, "fileName": file_name, "sheets": list(SHEETS)},
    }
    return {"events": [event]}


@router.get("/files/{fileName}")
def get_file(fileName: str):
    if "/" in fileName or "\\" in fileName or ".." in fileName or not SAFE_NAME.match(fileName):
        raise ApiError("not_found", "There is no workbook with that name.")
    path = resolve_file(fileName)
    if path is None:
        raise ApiError("not_found", "There is no workbook with that name.")
    try:
        stat = os.stat(path)
    except OSError:
        raise ApiError("not_found", "There is no workbook with that name.")
    return FileResponse(path, media_type=XLSX_TYPE, filename=fileName, stat_result=stat)
